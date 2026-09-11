import os
from openpyxl import Workbook, load_workbook

from path_utils import caminho_arquivo_local

ARQUIVO_TEMPLATES = caminho_arquivo_local("templates.xlsx")

TEMPLATE_PADRAO = "[cupom]"

# Apenas 2 templates predefinidos: produto e preço | produto, preço e cupom[cite: 1]
TEMPLATES_INICIAIS = {
    "Produto e Preço": "🔥 {nome}\n\n💰 Por apenas R$ {preco}\n\n🛒 Compre aqui:\n{link}",
    "Produto, Preço e Cupom": "🔥 {nome}\n\n💰 Por apenas R$ {preco}\n\n🎟️ Cupom: {cupom_usado}\n\n🛒 Compre aqui:\n{link}"
}


def inicializar_planilha_templates():
    """Garante que o arquivo Excel de templates exista e tenha dados iniciais."""
    if not os.path.exists(ARQUIVO_TEMPLATES):
        wb = Workbook()
        ws = wb.active
        if ws is not None:
            ws.title = "Templates"
            ws.append(["Nome", "Estrutura"])
            for nome, estrutura in TEMPLATES_INICIAIS.items():
                ws.append([nome, estrutura])
        wb.save(ARQUIVO_TEMPLATES)


def carregar_templates_do_excel():
    """Carrega os templates cadastrados na planilha Excel."""
    inicializar_planilha_templates()
    templates_dict = {}
    try:
        wb = load_workbook(ARQUIVO_TEMPLATES)
        ws = wb.active
        if ws is not None:
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] and row[1]:
                    templates_dict[str(row[0])] = str(row[1])
    except Exception:
        templates_dict = TEMPLATES_INICIAIS.copy()

    if not templates_dict:
        templates_dict = TEMPLATES_INICIAIS.copy()
    return templates_dict


def salvar_template_no_excel(nome, estrutura):
    """Adiciona ou atualiza um template na planilha Excel."""
    inicializar_planilha_templates()
    wb = load_workbook(ARQUIVO_TEMPLATES)
    ws = wb.active

    if ws is not None:
        encontrado = False
        for row in range(2, ws.max_row + 1):
            cell_val = ws.cell(row=row, column=1).value
            if cell_val == nome:
                ws.cell(row=row, column=2, value=estrutura)
                encontrado = True
                break

        if not encontrado:
            ws.append([nome, estrutura])

        wb.save(ARQUIVO_TEMPLATES)


def deletar_template_do_excel(nome):
    """Remove um template da planilha Excel pelo nome."""
    if not os.path.exists(ARQUIVO_TEMPLATES):
        return False
    wb = load_workbook(ARQUIVO_TEMPLATES)
    ws = wb.active

    if ws is not None:
        linha_encontrada = None
        for row in range(2, ws.max_row + 1):
            cell_val = ws.cell(row=row, column=1).value
            if cell_val == nome:
                linha_encontrada = row
                break

        if linha_encontrada:
            ws.delete_rows(linha_encontrada)
            wb.save(ARQUIVO_TEMPLATES)
            return True
    return False


def gerar_template(estilo, nome, preco, cupom, link, dicionario_templates=None):
    if not dicionario_templates:
        dicionario_templates = carregar_templates_do_excel()

    cupom_usado = cupom if cupom else TEMPLATE_PADRAO

    if estilo in dicionario_templates:
        estrutura = dicionario_templates[estilo]
        try:
            return estrutura.format(
                nome=nome,
                preco=preco,
                cupom=cupom,
                link=link,
                cupom_usado=cupom_usado
            )
        except Exception:
            pass

    return f"🔥 {nome}\n\n💵 R$ {preco}\n\n{cupom_usado}\n\n📎 Link:\n{link}"