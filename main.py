import json
import os
import sys
import tempfile
import threading
from datetime import datetime, timedelta
from io import BytesIO
from typing import Optional, cast
from tkinter import filedialog, messagebox
from urllib.request import urlopen

import customtkinter as ctk
import pyautogui
from PIL import Image, ImageTk
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill, Alignment

from path_utils import caminho_arquivo_local, caminho_recurso

# Módulos do projeto
from coordenadas import carregar_coordenadas, salvar_coordenadas
from templates import (
    carregar_templates_do_excel,
    salvar_template_no_excel,
    deletar_template_do_excel,
    TEMPLATE_PADRAO,
    gerar_template
)
from verificacao_imagem import eh_link_imagem_valido
from whatsapp import enviar_para_grupo


PASTA_TEMPORARIA = tempfile.gettempdir()
ARQUIVO_PRESETS = caminho_arquivo_local("coordenadas_presets.json")
ARQUIVO_TERMO = caminho_arquivo_local("termo_aceito.json")
ARQUIVO_EXCEL = caminho_arquivo_local("produtos.xlsx")

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

COR_JANELA = "#000000"
COR_CARD = "#262626"
COR_CAMPO = "#1a1a1a"
COR_BORDA = "#3a3a3a"
COR_TEXTO = "#ffffff"
COR_TEXTO_CINZA = "gray"
COR_BOTAO = "#3a3a3a"
COR_BOTAO_HOVER = "#4a4a4a"
COR_PRIMARIO = "#2f6fed"
COR_PRIMARIO_HOVER = "#255cc4"
COR_PERIGO = "#8a3a3a"
COR_PERIGO_HOVER = "#a94848"
COR_SUCESSO = "#2ecc71"
COR_AVISO = "#f39c12"
COR_WHATSAPP = "#25D366"
COR_QUADRADO_ICONE = "#3a3a3a"

FONTE_TITULO = ("Segoe UI", 16, "bold")
FONTE_LABEL = ("Segoe UI", 12)
FONTE_CAMPO = ("Segoe UI", 12)


def carregar_presets_coordenadas():
    if os.path.exists(ARQUIVO_PRESETS):
        try:
            with open(ARQUIVO_PRESETS, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {
        "Definição 1": {"grupo": {"x": 0, "y": 0}, "mensagem": {"x": 0, "y": 0}, "anexo": {"x": 0, "y": 0}, "fotos": {"x": 0, "y": 0}},
        "Definição 2": {"grupo": {"x": 0, "y": 0}, "mensagem": {"x": 0, "y": 0}, "anexo": {"x": 0, "y": 0}, "fotos": {"x": 0, "y": 0}},
        "Definição 3": {"grupo": {"x": 0, "y": 0}, "mensagem": {"x": 0, "y": 0}, "anexo": {"x": 0, "y": 0}, "fotos": {"x": 0, "y": 0}},
    }


def salvar_presets_coordenadas(presets_dict):
    try:
        with open(ARQUIVO_PRESETS, "w", encoding="utf-8") as f:
            json.dump(presets_dict, f, indent=4, ensure_ascii=False)
    except Exception as e:
        print(f"Erro ao salvar presets: {e}")


class AplicacaoAfiliado(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Automação Afiliados")
        self.geometry("1180x950")
        self.minsize(450, 500)
        self.configure(fg_color=COR_JANELA)
        self._aplicar_icone_janela()

        self.update_idletasks()
        self.after(100, self._verificar_termo_servico)

        self.imagem_atual = None
        self.caminho_imagem_local = None

        self.templates_dict = carregar_templates_do_excel()
        self.lista_templates = list(self.templates_dict.keys())
        self.indice_template_atual = 0

        self.agendamentos = []
        self.cancel_event = threading.Event()

        self._carregar_recursos()

        self.rodape = ctk.CTkFrame(
            self, fg_color="#121212", height=28, corner_radius=0)
        self.rodape.pack(side="bottom", fill="x")

        lbl_marca_dagua = ctk.CTkLabel(
            self.rodape,
            text="Desenvolvido por PietroHansel",
            font=("Segoe UI", 10, "bold"),
            text_color="#777777"
        )
        lbl_marca_dagua.pack(side="right", padx=15, pady=3)

        self.container_principal = ctk.CTkScrollableFrame(
            self,
            fg_color="transparent",
            scrollbar_fg_color="transparent",
            scrollbar_button_color="#333333",
            scrollbar_button_hover_color="#555555"
        )
        self.container_principal.pack(fill="both", expand=True)
        self._aplicar_efeito_autohide_scrollbar(
            self.container_principal, COR_JANELA)

        self.card_esquerdo: Optional[ctk.CTkFrame] = None
        self.card_direito: Optional[ctk.CTkFrame] = None
        self.modo_empilhado: Optional[bool] = None

        self._montar_painel_esquerdo()
        self._montar_painel_direito()

        self.bind("<Configure>", self._ao_redimensionar)

        self.update_idletasks()
        self._atualizar_layout_paineis(self.winfo_width())

    def _aplicar_icone_janela(self):
        try:
            caminho_icon = caminho_recurso("icone.png")
            if not os.path.exists(caminho_icon):
                caminho_icon = caminho_recurso("whatsapp.png")
            if not os.path.exists(caminho_icon):
                return

            imagem_icon = Image.open(caminho_icon).convert("RGBA")
            self._icone_janela = ImageTk.PhotoImage(imagem_icon)
            self.iconphoto(True, self._icone_janela)
        except Exception:
            pass

    def _verificar_termo_servico(self):
        """Exibe o Termo de Serviço ao abrir o software e exige autorização."""
        if os.path.exists(ARQUIVO_TERMO):
            try:
                with open(ARQUIVO_TERMO, "r", encoding="utf-8") as f:
                    dados = json.load(f)
                    if dados.get("aceito"):
                        return
            except Exception:
                pass

        janela_termo = ctk.CTkToplevel(self)
        janela_termo.title("Termo de Serviço e Uso")
        janela_termo.geometry("520x530")
        janela_termo.resizable(False, False)
        janela_termo.configure(fg_color=COR_JANELA)
        janela_termo.grab_set()

        ctk.CTkLabel(janela_termo, text="Termo de Serviço e Responsabilidade",
                     font=FONTE_TITULO, text_color="white").pack(pady=(15, 5))

        txt_termo = (
            "TERMOS DE USO E ISENÇÃO DE RESPONSABILIDADE\n\n"
            "1. ACEITAÇÃO DOS TERMOS\n"
            "Ao utilizar este software, você declara ter lido e concordado com todos os termos descritos abaixo.\n\n"
            "2. FINALIDADE E RESPONSABILIDADE DE USO\n"
            "Este software é uma ferramenta de automação para auxílio na divulgação de produtos e links de afiliados. "
            "O usuário é integralmente responsável por todas as mensagens, links e arquivos enviados através do sistema.\n\n"
            "3. ISENÇÃO DE RESPONSABILIDADE SOBRE BLOQUEIOS E BANIMENTOS\n"
            "O desenvolvedor (PietroHansel) não possui qualquer vínculo com as plataformas de terceiros (WhatsApp, Telegram, Shopee, Instagram, TikTok, etc.) "
            "e NÃO se responsabiliza por eventuais bloqueios, banimentos, restrições ou perdas decorrentes do envio de mensagens, uso de imagens, links ou prática de SPAM.\n\n"
            "4. MARCAS, ÍCONES E LOGOTIPOS\n"
            "Os ícones, logotipos, nomes e imagens de terceiros usados neste software (como WhatsApp, Instagram, TikTok, Telegram, Shopee e outros) são propriedade de seus respectivos detentores. "
            "Sua utilização neste aplicativo tem finalidade meramente identificadora e informativa, sem qualquer vínculo de patrocínio, aprovação, endosso ou associação por parte do desenvolvedor. "
            "O usuário deve respeitar os termos de uso das plataformas envolvidas e não deve reutilizar ou distribuir esses materiais sem autorização expressa dos seus titulares.\n\n"
            "5. PRIVACIDADE E DADOS\n"
            "As configurações, coordenadas capturadas e registros internos são salvos unicamente em seu computador local. "
            "O desenvolvedor não coleta, armazena ou transmite automaticamente dados pessoais do usuário, salvo quando o próprio usuário optar por compartilhá-los em contextos externos ao software.\n\n"
            "6. LIMITAÇÃO DE RESPONSABILIDADE\n"
            "O desenvolvedor não se responsabiliza por problemas de compatibilidade, falhas de terceiros, indisponibilidade de serviços externos, alterações na interface ou regras das plataformas citadas. "
            "O uso do software ocorre por conta e risco do usuário."
        )

        frame_txt = ctk.CTkFrame(
            janela_termo, fg_color=COR_CARD, corner_radius=8)
        frame_txt.pack(fill="both", expand=True, padx=15, pady=10)

        textbox = ctk.CTkTextbox(frame_txt, fg_color=COR_CAMPO,
                                 text_color=COR_TEXTO, font=("Segoe UI", 11), wrap="word")
        textbox.pack(fill="both", expand=True, padx=10, pady=10)
        textbox.insert("1.0", txt_termo)
        textbox.configure(state="disabled")

        termo_aceito_var = ctk.BooleanVar(value=False)

        def aceitar():
            try:
                with open(ARQUIVO_TERMO, "w", encoding="utf-8") as f:
                    json.dump({"aceito": True, "data": datetime.now().strftime(
                        "%Y-%m-%d %H:%M:%S")}, f, indent=4)
            except Exception as e:
                print(f"Erro ao salvar aceite do termo: {e}")
            termo_aceito_var.set(True)
            janela_termo.destroy()

        def recusar():
            try:
                janela_termo.destroy()
            except Exception:
                pass

            try:
                if self.winfo_exists():
                    self.destroy()
            except Exception:
                pass

            sys.exit()

        janela_termo.protocol("WM_DELETE_WINDOW", recusar)

        f_botoes = ctk.CTkFrame(janela_termo, fg_color="transparent")
        f_botoes.pack(fill="x", padx=15, pady=(0, 15))

        btn_recusar = ctk.CTkButton(f_botoes, text="Recusar e Sair", command=recusar,
                                    fg_color=COR_PERIGO, hover_color=COR_PERIGO_HOVER, width=140)
        btn_recusar.pack(side="left", padx=5)

        btn_aceitar = ctk.CTkButton(f_botoes, text="Concordar e Continuar",
                                    command=aceitar, fg_color=COR_SUCESSO, hover_color="#27ae60", width=180)
        btn_aceitar.pack(side="right", padx=5)

        self.wait_window(janela_termo)
        if not termo_aceito_var.get():
            try:
                if self.winfo_exists():
                    self.destroy()
            except Exception:
                pass
            sys.exit()

    def _ao_redimensionar(self, event):
        if event.widget != self:
            return
        self._atualizar_layout_paineis(event.width)

    def _atualizar_layout_paineis(self, largura):
        if self.card_esquerdo is None or self.card_direito is None:
            return

        e_empilhado = largura < 850

        if e_empilhado != self.modo_empilhado:
            self.modo_empilhado = e_empilhado
            if e_empilhado:
                self.container_principal.grid_columnconfigure(0, weight=1)
                self.container_principal.grid_columnconfigure(1, weight=0)
                self.card_esquerdo.grid(
                    row=0, column=0, padx=12, pady=(12, 6), sticky="ew")
                self.card_direito.grid(
                    row=1, column=0, padx=12, pady=(6, 12), sticky="ew")
            else:
                self.container_principal.grid_columnconfigure(0, weight=11)
                self.container_principal.grid_columnconfigure(1, weight=10)
                self.card_esquerdo.grid(row=0, column=0, padx=(
                    18, 9), pady=18, sticky="nsew")
                self.card_direito.grid(row=0, column=1, padx=(
                    9, 18), pady=18, sticky="nsew")

    def _carregar_recursos(self):
        def carregar_img(caminho, size):
            try:
                if os.path.exists(caminho):
                    img = Image.open(caminho)
                    return ctk.CTkImage(light_image=img, dark_image=img, size=size)
            except Exception:
                pass
            return None

        self.icone_upload = carregar_img(caminho_recurso("arquivo.png"), (64, 64))
        self.icone_nome = carregar_img(caminho_recurso("produto.png"), (18, 18))
        self.icone_categoria = carregar_img(caminho_recurso("categoria.png"), (18, 18))
        self.icone_preco = carregar_img(caminho_recurso("preco.png"), (18, 18))
        self.icone_cupom = carregar_img(caminho_recurso("cupom.png"), (18, 18))
        self.icone_link = carregar_img(caminho_recurso("link.png"), (18, 18))
        self.icone_imagem_url = carregar_img(caminho_recurso("link.png"), (18, 18))
        self.icone_descricao = carregar_img(caminho_recurso("descricao.png"), (18, 18))
        self.icone_agendar = carregar_img(caminho_recurso("clock.png"), (18, 18))
        self.icone_destino = carregar_img(caminho_recurso("destino1.png"), (18, 18))

        self.logo_excel = carregar_img(caminho_recurso("salvar.png"), (24, 24))
        self.icone_editar_planilha = carregar_img(caminho_recurso("editar.png"), (24, 24))
        self.icone_ver_agendamentos = carregar_img(caminho_recurso("agendamento.png"), (24, 24))
        self.icone_parar_envios = carregar_img(caminho_recurso("parar_envios.png"), (24, 24))
        self.logo_whatsapp = carregar_img(caminho_recurso("whatsapp.png"), (24, 24))
        self.logo_telegram = carregar_img(caminho_recurso("telegram.png"), (24, 24))
        self.logo_instagram = carregar_img(caminho_recurso("instagram.png"), (24, 24))
        self.logo_tiktok = carregar_img(caminho_recurso("tiktok.png"), (24, 24))

    def _aplicar_efeito_autohide_scrollbar(self, scrollable_frame, cor_fundo):
        scrollable_frame.configure(
            scrollbar_button_color=cor_fundo,
            scrollbar_button_hover_color=cor_fundo
        )

        def ao_entrar(e):
            scrollable_frame.configure(
                scrollbar_button_color="#333333",
                scrollbar_button_hover_color="#555555"
            )

        def ao_sair(e):
            scrollable_frame.configure(
                scrollbar_button_color=cor_fundo,
                scrollbar_button_hover_color=cor_fundo
            )

        scrollable_frame.bind("<Enter>", ao_entrar)
        scrollable_frame.bind("<Leave>", ao_sair)

    def _titulo_secao(self, master, texto):
        ctk.CTkLabel(
            master, text=texto, font=FONTE_TITULO, text_color=COR_TEXTO, anchor="w"
        ).pack(fill="x", pady=(0, 10))

    def _campo_rotulado(self, master, rotulo, placeholder, imagem_icone=None, icone_fallback="📦"):
        wrapper = ctk.CTkFrame(master, fg_color="transparent")
        cabecalho = ctk.CTkFrame(wrapper, fg_color="transparent")
        cabecalho.pack(fill="x", pady=(0, 4))

        quadrado_icone = ctk.CTkFrame(
            cabecalho, width=22, height=22, fg_color=COR_QUADRADO_ICONE, corner_radius=4
        )
        quadrado_icone.pack(side="left", padx=(0, 6))
        quadrado_icone.pack_propagate(False)

        if imagem_icone:
            lbl_icone = ctk.CTkLabel(
                quadrado_icone, text="", image=imagem_icone)
        else:
            lbl_icone = ctk.CTkLabel(
                quadrado_icone, text=icone_fallback, font=("Segoe UI", 11))

        lbl_icone.place(relx=0.5, rely=0.5, anchor="center")
        ctk.CTkLabel(cabecalho, text=rotulo, font=FONTE_LABEL,
                     text_color=COR_TEXTO, anchor="w").pack(side="left")

        entrada = ctk.CTkEntry(
            wrapper, placeholder_text=placeholder, font=FONTE_CAMPO, fg_color=COR_CAMPO,
            border_color=COR_BORDA, border_width=1, corner_radius=6, text_color=COR_TEXTO, height=34
        )
        entrada.pack(fill="x")
        return wrapper, entrada

    def _botao_rede_social(self, master, texto, comando, imagem_icone=None, icone_fallback="🔗"):
        btn_container = ctk.CTkFrame(
            master, fg_color="transparent", border_width=0, corner_radius=6, height=36, cursor="hand2"
        )

        def ao_clicar(e):
            comando()

        btn_container.bind("<Button-1>", ao_clicar)

        quadrado = ctk.CTkFrame(
            btn_container, width=22, height=22, fg_color=COR_QUADRADO_ICONE, corner_radius=4)
        quadrado.pack(side="left", padx=(8, 6), pady=7)
        quadrado.pack_propagate(False)
        quadrado.bind("<Button-1>", ao_clicar)

        if imagem_icone:
            lbl_icon = ctk.CTkLabel(quadrado, text="", image=imagem_icone)
        else:
            lbl_icon = ctk.CTkLabel(
                quadrado, text=icone_fallback, font=("Segoe UI", 11))

        lbl_icon.place(relx=0.5, rely=0.5, anchor="center")
        lbl_icon.bind("<Button-1>", ao_clicar)

        lbl_texto = ctk.CTkLabel(
            btn_container, text=texto, font=FONTE_LABEL, text_color=COR_TEXTO)
        lbl_texto.pack(side="left", padx=(0, 10))
        lbl_texto.bind("<Button-1>", ao_clicar)

        def on_enter(e):
            btn_container.configure(fg_color=COR_CAMPO)

        def on_leave(e):
            btn_container.configure(fg_color="transparent")

        btn_container.bind("<Enter>", on_enter)
        btn_container.bind("<Leave>", on_leave)

        return btn_container

    def _botao(self, master, texto, comando, cor=COR_BOTAO, cor_hover=COR_BOTAO_HOVER, largura=120, icone=None):
        return ctk.CTkButton(
            master, text=texto, image=icone, compound="left", command=comando, fg_color=cor,
            hover_color=cor_hover, text_color=COR_TEXTO, corner_radius=6, font=FONTE_LABEL,
            width=largura, height=34
        )

    def mostrar_feedback(self, texto, cor=COR_AVISO):
        self.resultado_status.configure(text=texto, text_color=cor)

    def _montar_painel_esquerdo(self):
        self.card_esquerdo = ctk.CTkFrame(
            self.container_principal, fg_color=COR_CARD, corner_radius=10)

        conteudo = ctk.CTkFrame(self.card_esquerdo, fg_color="transparent")
        conteudo.pack(fill="both", expand=True, padx=20, pady=20)

        self._titulo_secao(conteudo, "Dados do Produto")

        linha1 = ctk.CTkFrame(conteudo, fg_color="transparent")
        linha1.pack(fill="x", pady=(0, 12))
        linha1.grid_columnconfigure(0, weight=1)
        linha1.grid_columnconfigure(1, weight=1)

        wrap, self.campo_nome = self._campo_rotulado(
            linha1, "Nome do produto", "Ex: Tênis Fila Fitness", self.icone_nome, "🏷️")
        wrap.grid(row=0, column=0, sticky="ew", padx=(0, 4))

        wrap, self.campo_categoria = self._campo_rotulado(
            linha1, "Categoria", "Ex: Calçados", self.icone_categoria, "📂")
        wrap.grid(row=0, column=1, sticky="ew", padx=(4, 0))

        linha2 = ctk.CTkFrame(conteudo, fg_color="transparent")
        linha2.pack(fill="x", pady=(0, 12))
        linha2.grid_columnconfigure(0, weight=1)
        linha2.grid_columnconfigure(1, weight=1)

        wrap, self.campo_preco = self._campo_rotulado(
            linha2, "Preço", "Ex: 139,90", self.icone_preco, "💵")
        wrap.grid(row=0, column=0, sticky="ew", padx=(0, 4))

        wrap, self.campo_cupom = self._campo_rotulado(
            linha2, "Cupom", "Ex: DESCONTO10", self.icone_cupom, "🎟️")
        wrap.grid(row=0, column=1, sticky="ew", padx=(4, 0))

        wrap, self.campo_link = self._campo_rotulado(
            conteudo, "Link do produto", "Ex: https://shopee.com.br/...", self.icone_link, "🔗")
        wrap.pack(fill="x", pady=(0, 12))

        wrap, self.campo_link_imagem = self._campo_rotulado(
            conteudo, "Link da imagem (URL)", "Ex: https://...", self.icone_imagem_url, "🖼️")
        wrap.pack(fill="x", pady=(0, 20))

        self.campo_nome.bind(
            "<KeyRelease>", self.atualizar_template_em_tempo_real)
        self.campo_categoria.bind(
            "<KeyRelease>", self.atualizar_template_em_tempo_real)

        def evento_preco(e):
            self.validar_preco(e)
            self.atualizar_template_em_tempo_real(e)
        self.campo_preco.bind("<KeyRelease>", evento_preco)
        self.campo_preco.bind("<FocusOut>", self.atualizar_preco_campo)

        def evento_cupom(e):
            self.converter_cupom_maiuscula(e)
            self.atualizar_template_em_tempo_real(e)
        self.campo_cupom.bind("<KeyRelease>", evento_cupom)

        def evento_link(e):
            self.atualizar_template_em_tempo_real(e)
        self.campo_link.bind("<KeyRelease>", evento_link)

        self.campo_link_imagem.bind(
            "<KeyRelease>", self.validar_link_imagem_tempo_real)

        self._titulo_secao(conteudo, "Modelo da Publicação")

        cabecalho_desc = ctk.CTkFrame(conteudo, fg_color="transparent")
        cabecalho_desc.pack(fill="x", pady=(0, 4))

        quadrado_desc = ctk.CTkFrame(
            cabecalho_desc, width=22, height=22, fg_color=COR_QUADRADO_ICONE, corner_radius=4)
        quadrado_desc.pack(side="left", padx=(0, 6))
        quadrado_desc.pack_propagate(False)

        if self.icone_descricao:
            lbl_desc_icon = ctk.CTkLabel(
                quadrado_desc, text="", image=self.icone_descricao)
        else:
            lbl_desc_icon = ctk.CTkLabel(
                quadrado_desc, text="📝", font=("Segoe UI", 11))

        lbl_desc_icon.place(relx=0.5, rely=0.5, anchor="center")
        ctk.CTkLabel(cabecalho_desc, text="Descrição", font=FONTE_LABEL,
                     text_color=COR_TEXTO, anchor="w").pack(side="left")

        self.campo_descricao = ctk.CTkTextbox(
            conteudo, height=90, fg_color=COR_CAMPO, border_color=COR_BORDA, border_width=1,
            corner_radius=6, text_color=COR_TEXTO, font=FONTE_CAMPO
        )
        self.campo_descricao.pack(fill="x", pady=(0, 10))

        placeholder_desc_inicial = "Escreva ou aplique um template para gerar a descrição..."
        self.campo_descricao.insert("1.0", placeholder_desc_inicial)
        self.campo_descricao.configure(text_color=COR_TEXTO_CINZA)

        def ao_focar_descricao(e):
            if self.campo_descricao.get("1.0", "end-1c") == placeholder_desc_inicial:
                self.campo_descricao.delete("1.0", "end")
                self.campo_descricao.configure(text_color=COR_TEXTO)

        def ao_perder_foco_descricao(e):
            if not self.campo_descricao.get("1.0", "end-1c").strip():
                self.campo_descricao.insert("1.0", placeholder_desc_inicial)
                self.campo_descricao.configure(text_color=COR_TEXTO_CINZA)

        self.campo_descricao.bind("<FocusIn>", ao_focar_descricao)
        self.campo_descricao.bind("<FocusOut>", ao_perder_foco_descricao)
        self.campo_descricao.bind(
            "<KeyRelease>", self.atualizar_descricao_preview)

        linha_botoes_template = ctk.CTkFrame(conteudo, fg_color="transparent")
        linha_botoes_template.pack(fill="x", pady=(0, 20))

        self.botao_template_descricao = self._botao(
            linha_botoes_template, "Aplicar template (nenhum)", self.aplicar_template_descricao, largura=160
        )
        self.botao_template_descricao.pack(side="left", padx=(0, 6))

        self._botao(
            linha_botoes_template, "Criar", self.abrir_janela_criar_template, cor=COR_BOTAO, cor_hover=COR_SUCESSO, largura=65
        ).pack(side="left", padx=(0, 6))

        self._botao(
            linha_botoes_template, "Editar", self.abrir_janela_editar_template, cor=COR_BOTAO, cor_hover=COR_PRIMARIO_HOVER, largura=65
        ).pack(side="left", padx=(0, 6))

        self._botao(
            linha_botoes_template, "Deletar", self.abrir_janela_deletar_template, cor=COR_BOTAO, cor_hover=COR_PERIGO, largura=65
        ).pack(side="left")

        self._titulo_secao(conteudo, "Cadastro e Controle")
        linha_controle = ctk.CTkFrame(conteudo, fg_color="transparent")
        linha_controle.pack(fill="x", pady=(0, 20))

        linha_controle.grid_columnconfigure(0, weight=1)
        linha_controle.grid_columnconfigure(1, weight=1)

        b_salvar = self._botao_rede_social(
            linha_controle, "Salvar produto", self.salvar_no_excel, imagem_icone=self.logo_excel)
        b_salvar.grid(row=0, column=0, sticky="ew", padx=2, pady=2)

        b_editar = self._botao_rede_social(linha_controle, "Editar planilha", self.abrir_janela_edicao_planilha,
                                           imagem_icone=self.icone_editar_planilha, icone_fallback="📊")
        b_editar.grid(row=0, column=1, sticky="ew", padx=2, pady=2)

        b_agend = self._botao_rede_social(linha_controle, "Ver Agendamentos", self.abrir_janela_agendamentos,
                                          imagem_icone=self.icone_ver_agendamentos, icone_fallback="📅")
        b_agend.grid(row=1, column=0, sticky="ew", padx=2, pady=2)

        b_parar = self._botao_rede_social(
            linha_controle, "Parar Envios", self.parar_envios, imagem_icone=self.icone_parar_envios, icone_fallback="🛑")
        b_parar.grid(row=1, column=1, sticky="ew", padx=2, pady=2)

        self._titulo_secao(conteudo, "Formas de Envio")
        linha_canais = ctk.CTkFrame(conteudo, fg_color="transparent")
        linha_canais.pack(fill="x", pady=(0, 10))

        linha_canais.grid_columnconfigure(0, weight=1)
        linha_canais.grid_columnconfigure(1, weight=1)

        b_wa = self._botao_rede_social(
            linha_canais, "WhatsApp", self.enviar_para_whatsapp_grupo, imagem_icone=self.logo_whatsapp)
        b_wa.grid(row=0, column=0, sticky="ew", padx=2, pady=2)

        b_tg = self._botao_rede_social(linha_canais, "Telegram", lambda: self.mostrar_feedback(
            "Integração Telegram em breve!", COR_AVISO), imagem_icone=self.logo_telegram, icone_fallback="✈️")
        b_tg.grid(row=0, column=1, sticky="ew", padx=2, pady=2)

        b_ig = self._botao_rede_social(linha_canais, "Instagram", lambda: self.mostrar_feedback(
            "Integração Instagram em breve!", COR_AVISO), imagem_icone=self.logo_instagram, icone_fallback="📸")
        b_ig.grid(row=1, column=0, sticky="ew", padx=2, pady=2)

        b_tk = self._botao_rede_social(linha_canais, "TikTok", lambda: self.mostrar_feedback(
            "Integração TikTok em breve!", COR_AVISO), imagem_icone=self.logo_tiktok, icone_fallback="🎵")
        b_tk.grid(row=1, column=1, sticky="ew", padx=2, pady=2)

        self.resultado_status = ctk.CTkLabel(
            conteudo, text="", font=FONTE_LABEL, text_color="gray")
        self.resultado_status.pack(fill="x", pady=(15, 0))

    def _montar_painel_direito(self):
        self.card_direito = ctk.CTkFrame(
            self.container_principal, fg_color=COR_CARD, corner_radius=10)

        ctk.CTkLabel(self.card_direito, text="Visualização da Publicação",
                     font=FONTE_TITULO, text_color=COR_TEXTO).pack(pady=(20, 10))

        container_preview = ctk.CTkFrame(
            self.card_direito, fg_color="transparent")
        container_preview.pack(fill="both", expand=True, padx=20, pady=(0, 20))

        container_preview.grid_columnconfigure(0, weight=1)
        container_preview.grid_rowconfigure(0, weight=1)
        container_preview.grid_rowconfigure(1, weight=1)

        self.area_preview = ctk.CTkFrame(
            container_preview, fg_color=COR_CAMPO, corner_radius=8, border_color=COR_BORDA, border_width=1, cursor="hand2", height=240
        )
        self.area_preview.grid(row=0, column=0, sticky="nsew", pady=(0, 6))
        self.area_preview.bind(
            "<Button-1>", lambda e: self.escolher_imagem_local())

        self.rotulo_imagem = ctk.CTkLabel(
            self.area_preview, text="Clique para carregar arquivo local\nou insira uma URL ao lado",
            font=("Segoe UI", 12), text_color="gray", image=self.icone_upload, compound="top"
        )
        self.rotulo_imagem.place(relx=0.5, rely=0.5, anchor="center")
        self.rotulo_imagem.bind(
            "<Button-1>", lambda e: self.escolher_imagem_local())

        self.container_descricao_preview = ctk.CTkFrame(
            container_preview, fg_color=COR_CAMPO, corner_radius=8, border_color=COR_BORDA, border_width=1
        )
        self.container_descricao_preview.grid(
            row=1, column=0, sticky="nsew", pady=(6, 0))

        self.scroll_descricao = ctk.CTkScrollableFrame(
            self.container_descricao_preview,
            fg_color="transparent",
            scrollbar_fg_color="transparent",
            scrollbar_button_color="#333333",
            scrollbar_button_hover_color="#555555"
        )
        self.scroll_descricao.pack(fill="both", expand=True, padx=10, pady=10)
        self._aplicar_efeito_autohide_scrollbar(
            self.scroll_descricao, COR_CAMPO)

        self.previa_descricao = ctk.CTkLabel(
            self.scroll_descricao, text="A visualização da descrição aparecerá aqui...", font=("Segoe UI", 11),
            text_color="gray", wraplength=340, anchor="nw", justify="left"
        )
        self.previa_descricao.pack(fill="both", expand=True)

        def _ajustar_quebra_texto(event):
            largura_util = event.width - 25
            if largura_util > 50:
                self.previa_descricao.configure(wraplength=largura_util)

        self.scroll_descricao.bind("<Configure>", _ajustar_quebra_texto)

    def parar_envios(self):
        self.cancel_event.set()

        cancelados = 0
        for item in self.agendamentos:
            if item["status"] == "Agendado":
                if item.get("timer"):
                    item["timer"].cancel()
                item["status"] = "Cancelado"
                cancelados += 1

        self.mostrar_feedback(
            "Envios parados/cancelados pelo usuário!", COR_PERIGO)
        messagebox.showinfo(
            "Envios Parados", f"Envios em andamento interrompidos. {cancelados} agendamentos cancelados.")

    def abrir_janela_agendamentos(self):
        janela = ctk.CTkToplevel(self)
        janela.title("Agendamentos de Envio")
        janela.geometry("680x480")
        janela.minsize(400, 350)
        janela.configure(fg_color=COR_JANELA)
        janela.update()
        janela.grab_set()

        ctk.CTkLabel(janela, text="Fila de Envios Agendados",
                     font=FONTE_TITULO, text_color="white").pack(pady=15)

        frame_lista = ctk.CTkScrollableFrame(
            janela, fg_color=COR_CARD, scrollbar_fg_color="transparent",
            scrollbar_button_color="#333333", scrollbar_button_hover_color="#555555"
        )
        frame_lista.pack(fill="both", expand=True, padx=20, pady=(0, 15))
        self._aplicar_efeito_autohide_scrollbar(frame_lista, COR_CARD)

        def atualizar_lista_agendamentos():
            for widget in frame_lista.winfo_children():
                widget.destroy()

            if not self.agendamentos:
                ctk.CTkLabel(
                    frame_lista, text="Nenhum agendamento pendente.", text_color="gray").pack(pady=20)
                return

            for ag in list(self.agendamentos):
                f_row = ctk.CTkFrame(
                    frame_lista, fg_color=COR_CAMPO, corner_radius=6)
                f_row.pack(fill="x", pady=4, padx=5)

                qtd_prods = len(ag["produtos"])
                info = f"#{ag['id']} | Horário: {ag['horario_str']} | Grupo: {ag['grupo']} | {qtd_prods} produto(s) | Status: {ag['status']}"

                cor_status = COR_SUCESSO if ag['status'] == 'Concluído' else (
                    COR_PERIGO if ag['status'] == 'Cancelado' else COR_AVISO)

                ctk.CTkLabel(f_row, text=info, font=FONTE_LABEL, text_color=cor_status, anchor="w").pack(
                    side="left", padx=10, pady=8)

                if ag["status"] == "Agendado":
                    def cancelar_este(ag_obj=ag):
                        if ag_obj.get("timer"):
                            ag_obj["timer"].cancel()
                        ag_obj["status"] = "Cancelado"
                        atualizar_lista_agendamentos()
                        self.mostrar_feedback(
                            f"Agendamento #{ag_obj['id']} cancelado!", COR_AVISO)

                    ctk.CTkButton(
                        f_row, text="Cancelar", fg_color=COR_PERIGO, hover_color=COR_PERIGO_HOVER,
                        width=80, height=28, command=cancelar_este
                    ).pack(side="right", padx=10)

        atualizar_lista_agendamentos()

    def formatar_preco(self, preco_str):
        if not preco_str:
            return "[Preco]"
        preco_str = str(preco_str).strip()
        if "," not in preco_str and "." not in preco_str:
            preco_str += ",00"
        return preco_str

    def validar_preco(self, event=None):
        preco_atual = self.campo_preco.get()
        preco_filtrado = "".join(
            c for c in preco_atual if c.isdigit() or c in ".,")
        if preco_atual != preco_filtrado:
            self.campo_preco.delete(0, "end")
            self.campo_preco.insert(0, preco_filtrado)
        self.atualizar_template_em_tempo_real()

    def atualizar_preco_campo(self, event=None):
        preco_atual = self.campo_preco.get()
        preco_formatado = self.formatar_preco(preco_atual)
        if preco_atual != preco_formatado:
            self.campo_preco.delete(0, "end")
            self.campo_preco.insert(0, preco_formatado)
        self.atualizar_template_em_tempo_real()

    def formatar_horario(self, horario_str):
        if horario_str is None:
            return ""

        texto = str(horario_str).strip()
        if not texto:
            return ""

        # Mantém somente dígitos e separadores válidos, evitando duplicações
        # como "143014:30:00" ao limpar ou editar o campo.
        texto = "".join(c for c in texto if c.isdigit() or c == ":")
        texto = texto.strip(":")

        if not texto:
            return ""

        # Remove separadores extras e recria o formato em um único padrão.
        digitos = "".join(c for c in texto if c.isdigit())[:6]

        if len(digitos) <= 2:
            return digitos
        if len(digitos) <= 4:
            return f"{digitos[:2]}:{digitos[2:]}"
        return f"{digitos[:2]}:{digitos[2:4]}:{digitos[4:]}"

    def validar_horario(self, event=None):
        if event and hasattr(event, "widget"):
            entry = event.widget
            horario_atual = entry.get()
            horario_filtrado = "".join(
                c for c in horario_atual if c.isdigit() or c == ":")
            if horario_atual != horario_filtrado:
                entry.delete(0, "end")
                entry.insert(0, horario_filtrado)

    def atualizar_horario_campo(self, event=None):
        if event and hasattr(event, "widget"):
            entry = event.widget
            horario_atual = entry.get()
            horario_formatado = self.formatar_horario(horario_atual)
            if horario_atual != horario_formatado:
                entry.delete(0, "end")
                entry.insert(0, horario_formatado)

    def converter_cupom_maiuscula(self, event=None):
        cupom_atual = self.campo_cupom.get()
        cupom_maiuscula = cupom_atual.upper()
        if cupom_atual != cupom_maiuscula:
            self.campo_cupom.delete(0, "end")
            self.campo_cupom.insert(0, cupom_maiuscula)
        self.atualizar_template_em_tempo_real()

    def validar_link_imagem_tempo_real(self, event=None):
        link = self.campo_link_imagem.get().strip()
        if not link:
            self.campo_link_imagem.configure(border_color=COR_BORDA)
            self.limpar_imagem_preview_visual()
            return

        valido, msg = eh_link_imagem_valido(link)
        if valido:
            self.campo_link_imagem.configure(border_color=COR_SUCESSO)
            self.carregar_imagem_por_url(link)
        else:
            self.campo_link_imagem.configure(border_color=COR_PERIGO)

    def escolher_imagem_local(self):
        caminho = filedialog.askopenfilename(
            title="Selecione a imagem do produto",
            filetypes=[("Imagens", "*.png *.jpg *.jpeg *.webp *.gif")],
        )
        if not caminho:
            return
        try:
            img = Image.open(caminho)
            img.thumbnail((320, 320))
            self.imagem_atual = ctk.CTkImage(
                light_image=img, dark_image=img, size=img.size)
            self.rotulo_imagem.configure(image=self.imagem_atual, text="")
            self.caminho_imagem_local = caminho

            self.campo_link_imagem.delete(0, "end")
            self.campo_link_imagem.insert(0, caminho)
            self.campo_link_imagem.configure(border_color=COR_SUCESSO)

            self.mostrar_feedback("Imagem local carregada!", COR_SUCESSO)
        except Exception as e:
            messagebox.showerror("Erro ao carregar imagem", str(e))

    def carregar_imagem_por_url(self, url):
        try:
            req = urlopen(url, timeout=5)
            dados_img = req.read()
            img = Image.open(BytesIO(dados_img))
            img.thumbnail((320, 320))
            self.imagem_atual = ctk.CTkImage(
                light_image=img, dark_image=img, size=img.size)
            self.rotulo_imagem.configure(image=self.imagem_atual, text="")
            self.caminho_imagem_local = None
        except Exception:
            pass

    def limpar_imagem_preview_visual(self):
        self.imagem_atual = None
        self.caminho_imagem_local = None
        self.rotulo_imagem.configure(
            image=self.icone_upload, text="Clique para carregar arquivo local\nou insira uma URL ao lado")

    def atualizar_descricao_preview(self, event=None):
        descricao = self.campo_descricao.get("1.0", "end-1c")
        if descricao == "Escreva ou aplique um template para gerar a descrição...":
            self.previa_descricao.configure(
                text="A visualização da descrição aparecerá aqui...", text_color="gray")
        else:
            self.previa_descricao.configure(text=descricao if descricao.strip(
            ) else "A visualização da descrição aparecerá aqui...", text_color=COR_TEXTO)

    def atualizar_template_em_tempo_real(self, event=None):
        descricao_atual = self.campo_descricao.get("1.0", "end-1c")
        if descricao_atual == "Escreva ou aplique um template para gerar a descrição...":
            return

        if self.lista_templates:
            estilo_atual = self.lista_templates[self.indice_template_atual -
                                                1 if self.indice_template_atual > 0 else 0]
            nome = self.campo_nome.get() or "[Produto]"
            preco = self.formatar_preco(self.campo_preco.get())
            cupom = self.campo_cupom.get()
            link = self.campo_link.get() or "[Link do produto]"

            template = gerar_template(
                estilo_atual, nome, preco, cupom, link, self.templates_dict
            )

            self.campo_descricao.delete("1.0", "end")
            self.campo_descricao.insert("1.0", template)
            self.campo_descricao.configure(text_color=COR_TEXTO)
            self.atualizar_descricao_preview()

    def aplicar_template_descricao(self):
        self.templates_dict = carregar_templates_do_excel()
        self.lista_templates = list(self.templates_dict.keys())

        if not self.lista_templates:
            self.mostrar_feedback("Nenhum template disponível.", COR_AVISO)
            return

        if self.indice_template_atual >= len(self.lista_templates):
            self.indice_template_atual = 0

        estilo = self.lista_templates[self.indice_template_atual]
        nome = self.campo_nome.get() or "[Produto]"
        preco = self.formatar_preco(self.campo_preco.get())
        cupom = self.campo_cupom.get()
        link = self.campo_link.get() or "[Link do produto]"

        template = gerar_template(
            estilo, nome, preco, cupom, link, self.templates_dict)
        self.campo_descricao.delete("1.0", "end")
        self.campo_descricao.insert("1.0", template)
        self.campo_descricao.configure(text_color=COR_TEXTO)
        self.atualizar_descricao_preview()

        self.botao_template_descricao.configure(
            text=f"Aplicar template ({estilo})")

        self.indice_template_atual = (
            self.indice_template_atual + 1) % len(self.lista_templates)

    def abrir_janela_criar_template(self):
        janela = ctk.CTkToplevel(self)
        janela.title("Criar Novo Template")
        janela.geometry("480x420")
        janela.minsize(380, 350)
        janela.configure(fg_color=COR_JANELA)
        janela.update()
        janela.grab_set()

        ctk.CTkLabel(janela, text="Criar Novo Template na Planilha",
                     font=FONTE_TITULO, text_color="white").pack(pady=15)

        frame = ctk.CTkFrame(janela, fg_color=COR_CARD, corner_radius=10)
        frame.pack(fill="both", expand=True, padx=20, pady=10)

        ctk.CTkLabel(frame, text="Nome do Template", font=FONTE_LABEL,
                     text_color="gray").pack(anchor="w", padx=15, pady=(15, 2))
        campo_nome_template = ctk.CTkEntry(
            frame, placeholder_text="Ex: PromocaoRelampago", height=35, fg_color=COR_CAMPO, border_color=COR_BORDA)
        campo_nome_template.pack(fill="x", padx=15, pady=(0, 10))

        ctk.CTkLabel(frame, text="Estrutura (Use: {nome}, {preco}, {cupom}, {link})",
                     font=FONTE_LABEL, text_color="gray").pack(anchor="w", padx=15, pady=(5, 2))

        campo_texto_template = ctk.CTkTextbox(
            frame, height=140, fg_color=COR_CAMPO, border_color=COR_BORDA, border_width=1,
            corner_radius=6, text_color=COR_TEXTO, font=FONTE_CAMPO
        )
        campo_texto_template.pack(fill="x", padx=15, pady=(0, 15))

        exemplo_estrutura = "⚡ {nome} ⚡\n\nPor Apenas: R$ {preco}\n\n{cupom}\n\nGaranta o seu:\n{link}"
        campo_texto_template.insert("1.0", exemplo_estrutura)

        def salvar_novo_template():
            nome_tpl = campo_nome_template.get().strip()
            corpo_tpl = campo_texto_template.get("1.0", "end-1c").strip()

            if not nome_tpl or not corpo_tpl:
                messagebox.showerror(
                    "Erro", "Preencha o nome e a estrutura do template!")
                return

            salvar_template_no_excel(nome_tpl, corpo_tpl)
            self.templates_dict = carregar_templates_do_excel()
            self.lista_templates = list(self.templates_dict.keys())

            if nome_tpl in self.lista_templates:
                self.indice_template_atual = self.lista_templates.index(
                    nome_tpl)

            janela.destroy()
            self.mostrar_feedback(
                f"Template '{nome_tpl}' salvo na planilha com sucesso!", COR_SUCESSO)

        self._botao(janela, "Salvar na Planilha", salvar_novo_template,
                    cor=COR_BOTAO, cor_hover="#27ae60", largura=180).pack(pady=15)

    def abrir_janela_editar_template(self):
        self.templates_dict = carregar_templates_do_excel()
        self.lista_templates = list(self.templates_dict.keys())

        if not self.lista_templates:
            self.mostrar_feedback(
                "Não há templates disponíveis para editar.", COR_AVISO)
            return

        janela = ctk.CTkToplevel(self)
        janela.title("Editar Template")
        janela.geometry("400x350")
        janela.minsize(320, 250)
        janela.configure(fg_color=COR_JANELA)
        janela.update()
        janela.grab_set()

        ctk.CTkLabel(janela, text="Selecione o Template para Editar",
                     font=FONTE_TITULO, text_color="white").pack(pady=15)

        frame_lista = ctk.CTkScrollableFrame(
            janela,
            fg_color=COR_CARD,
            scrollbar_fg_color="transparent",
            scrollbar_button_color="#333333",
            scrollbar_button_hover_color="#555555"
        )
        frame_lista.pack(fill="both", expand=True, padx=20, pady=10)
        self._aplicar_efeito_autohide_scrollbar(frame_lista, COR_CARD)

        for template_nome, template_corpo in self.templates_dict.items():
            f_row = ctk.CTkFrame(
                frame_lista, fg_color=COR_CAMPO, corner_radius=6)
            f_row.pack(fill="x", pady=4, padx=5)

            ctk.CTkLabel(f_row, text=template_nome, font=FONTE_LABEL,
                         text_color=COR_TEXTO, anchor="w").pack(side="left", padx=10, pady=8)

            def editar_este_template(nome=template_nome, corpo=template_corpo):
                janela.destroy()
                self._abrir_formulario_edicao_template(nome, corpo)

            ctk.CTkButton(
                f_row, text="Editar", fg_color=COR_BOTAO, hover_color=COR_PRIMARIO_HOVER,
                width=70, height=28, command=editar_este_template
            ).pack(side="right", padx=10)

    def _abrir_formulario_edicao_template(self, nome_tpl_original, corpo_tpl_original):
        janela = ctk.CTkToplevel(self)
        janela.title(f"Editar Template: {nome_tpl_original}")
        janela.geometry("480x420")
        janela.minsize(380, 350)
        janela.configure(fg_color=COR_JANELA)
        janela.update()
        janela.grab_set()

        ctk.CTkLabel(janela, text=f"Editar Template '{nome_tpl_original}'",
                     font=FONTE_TITULO, text_color="white").pack(pady=15)

        frame = ctk.CTkFrame(janela, fg_color=COR_CARD, corner_radius=10)
        frame.pack(fill="both", expand=True, padx=20, pady=10)

        ctk.CTkLabel(frame, text="Nome do Template", font=FONTE_LABEL,
                     text_color="gray").pack(anchor="w", padx=15, pady=(15, 2))
        campo_nome_template = ctk.CTkEntry(
            frame, height=35, fg_color=COR_CAMPO, border_color=COR_BORDA)
        campo_nome_template.insert(0, nome_tpl_original)
        campo_nome_template.pack(fill="x", padx=15, pady=(0, 10))

        ctk.CTkLabel(frame, text="Estrutura (Use: {nome}, {preco}, {cupom}, {link})",
                     font=FONTE_LABEL, text_color="gray").pack(anchor="w", padx=15, pady=(5, 2))

        campo_texto_template = ctk.CTkTextbox(
            frame, height=140, fg_color=COR_CAMPO, border_color=COR_BORDA, border_width=1,
            corner_radius=6, text_color=COR_TEXTO, font=FONTE_CAMPO
        )
        campo_texto_template.pack(fill="x", padx=15, pady=(0, 15))
        campo_texto_template.insert("1.0", corpo_tpl_original)

        def salvar_edicao():
            novo_nome = campo_nome_template.get().strip()
            novo_corpo = campo_texto_template.get("1.0", "end-1c").strip()

            if not novo_nome or not novo_corpo:
                messagebox.showerror(
                    "Erro", "Preencha o nome e a estrutura do template!")
                return

            if novo_nome != nome_tpl_original:
                deletar_template_do_excel(nome_tpl_original)

            salvar_template_no_excel(novo_nome, novo_corpo)
            self.templates_dict = carregar_templates_do_excel()
            self.lista_templates = list(self.templates_dict.keys())

            if novo_nome in self.lista_templates:
                self.indice_template_atual = self.lista_templates.index(
                    novo_nome)

            janela.destroy()
            self.mostrar_feedback(
                f"Template '{novo_nome}' atualizado com sucesso!", COR_SUCESSO)

        self._botao(janela, "Salvar Alterações", salvar_edicao,
                    cor=COR_SUCESSO, cor_hover="#27ae60", largura=180).pack(pady=15)

    def abrir_janela_deletar_template(self):
        self.templates_dict = carregar_templates_do_excel()
        self.lista_templates = list(self.templates_dict.keys())

        if not self.lista_templates:
            self.mostrar_feedback(
                "Não há templates disponíveis para deletar.", COR_AVISO)
            return

        janela = ctk.CTkToplevel(self)
        janela.title("Deletar Template")
        janela.geometry("400x350")
        janela.minsize(320, 250)
        janela.configure(fg_color=COR_JANELA)
        janela.update()
        janela.grab_set()

        ctk.CTkLabel(janela, text="Selecione o Template para Excluir",
                     font=FONTE_TITULO, text_color="white").pack(pady=15)

        frame_lista = ctk.CTkScrollableFrame(
            janela,
            fg_color=COR_CARD,
            scrollbar_fg_color="transparent",
            scrollbar_button_color="#333333",
            scrollbar_button_hover_color="#555555"
        )
        frame_lista.pack(fill="both", expand=True, padx=20, pady=10)
        self._aplicar_efeito_autohide_scrollbar(frame_lista, COR_CARD)

        def atualizar_lista_templates():
            for widget in frame_lista.winfo_children():
                widget.destroy()

            self.templates_dict = carregar_templates_do_excel()
            self.lista_templates = list(self.templates_dict.keys())

            if not self.lista_templates:
                janela.destroy()
                self.mostrar_feedback(
                    "Todos os templates foram removidos.", COR_AVISO)
                return

            for template_nome in list(self.lista_templates):
                f_row = ctk.CTkFrame(
                    frame_lista, fg_color=COR_CAMPO, corner_radius=6)
                f_row.pack(fill="x", pady=4, padx=5)

                ctk.CTkLabel(f_row, text=template_nome, font=FONTE_LABEL,
                             text_color=COR_TEXTO, anchor="w").pack(side="left", padx=10, pady=8)

                def excluir_template_por_nome(nome=template_nome):
                    if len(self.lista_templates) <= 1:
                        messagebox.showwarning(
                            "Aviso", "Você precisa manter pelo menos um template cadastrado!")
                        return

                    if messagebox.askyesno("Confirmar Exclusão", f"Deseja excluir o template '{nome}'?"):
                        deletar_template_do_excel(nome)
                        self.templates_dict = carregar_templates_do_excel()
                        self.lista_templates = list(self.templates_dict.keys())
                        self.indice_template_atual = 0
                        self.botao_template_descricao.configure(
                            text="Aplicar template (nenhum)")
                        atualizar_lista_templates()
                        self.mostrar_feedback(
                            f"Template '{nome}' excluído com sucesso!", COR_AVISO)

                ctk.CTkButton(
                    f_row, text="Deletar", fg_color=COR_BOTAO, hover_color=COR_PERIGO_HOVER,
                    width=70, height=28, command=excluir_template_por_nome
                ).pack(side="right", padx=10)

        atualizar_lista_templates()

    def preparar_planilha_excel(self, worksheet: Worksheet):
        if worksheet is None:
            return
        cabecalhos = ["ID", "Data", "Nome do Produto", "Preco", "Cupom",
                      "Link do Produto", "Link da Imagem", "Descricao", "Categoria"]
        if worksheet.max_row == 1 and worksheet["A1"].value is None:
            worksheet.append(cabecalhos)
        elif worksheet["A1"].value != "ID":
            worksheet.insert_cols(1)
            worksheet["A1"] = "ID"
            for col, val in enumerate(cabecalhos, start=1):
                cast(Cell, worksheet.cell(row=1, column=col)).value = val

        if worksheet.max_column < 9:
            cast(Cell, worksheet.cell(row=1, column=9)).value = "Categoria"

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

    def salvar_no_excel(self):
        nome = self.campo_nome.get().strip()
        preco = self.campo_preco.get().strip()
        link = self.campo_link.get().strip()
        link_imagem = self.campo_link_imagem.get().strip()
        descricao = self.campo_descricao.get("1.0", "end-1c")

        if not nome or not preco or not link or not link_imagem or descricao == "Escreva ou aplique um template para gerar a descrição...":
            self.mostrar_feedback(
                "Preencha todos os campos obrigatórios!", COR_PERIGO)
            messagebox.showwarning(
                "Campos Incompletos", "Por favor, preencha todos os dados e a descrição antes de salvar.")
            return

        try:
            if os.path.exists(ARQUIVO_EXCEL):
                workbook = load_workbook(ARQUIVO_EXCEL)
                worksheet = workbook.active
                if worksheet is None:
                    worksheet = workbook.create_sheet("Produtos")
            else:
                workbook = Workbook()
                worksheet = workbook.active
                if worksheet is not None:
                    worksheet.title = "Produtos"

            if worksheet is None:
                raise RuntimeError("Não foi possível obter a planilha ativa.")
            worksheet = cast(Worksheet, worksheet)
            self.preparar_planilha_excel(worksheet)

            produtos_existentes = self.carregar_produtos_excel()
            for p in produtos_existentes:
                if (p["link"].strip().lower() == link.lower()) or (p["nome"].strip().lower() == nome.lower() and str(p["preco"]).strip() == preco):
                    self.mostrar_feedback(
                        f"Produto duplicado! Já cadastrado (ID #{p['id']}).", COR_AVISO)
                    messagebox.showwarning(
                        "Produto Duplicado", f"Este produto já consta na planilha sob o ID #{p['id']}.")
                    return

            data = datetime.now().strftime("%d/%m/%Y %H:%M")
            cupom = self.campo_cupom.get().strip() or ""
            categoria = self.campo_categoria.get().strip() or ""

            ids = []
            if worksheet is not None:
                for row in worksheet.iter_rows(min_row=2, values_only=True):
                    if row and row[0] is not None:
                        try:
                            ids.append(int(float(str(row[0]))))
                        except ValueError:
                            pass

                proximo_id = max(ids, default=0) + 1
                worksheet.append([proximo_id, data, nome, preco,
                                 cupom, link, link_imagem, descricao, categoria])

                workbook.save(ARQUIVO_EXCEL)

                self.mostrar_feedback(
                    f"Produto #{proximo_id} salvo e planilha atualizada!", COR_SUCESSO)
        except Exception as e:
            self.mostrar_feedback(
                f"Erro ao salvar no Excel: {str(e)}", COR_PERIGO)

    def carregar_produtos_excel(self):
        if not os.path.exists(ARQUIVO_EXCEL):
            return []
        try:
            workbook = load_workbook(ARQUIVO_EXCEL)
            worksheet = workbook.active
            if worksheet is None:
                return []
            worksheet = cast(Worksheet, worksheet)
            produtos = []
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 8 and row[0] is not None:
                    try:
                        prod_id = int(float(str(row[0]).strip()))
                    except (ValueError, TypeError):
                        try:
                            prod_id = int(str(row[0]).strip())
                        except (ValueError, TypeError):
                            prod_id = 0

                    categoria_val = row[8] if len(
                        row) > 8 and row[8] is not None else ""

                    produtos.append({
                        "id": prod_id,
                        "data": row[1] or "",
                        "nome": row[2] or "",
                        "preco": row[3] or "",
                        "cupom": row[4] or "",
                        "link": row[5] or "",
                        "link_imagem": row[6] or "",
                        "descricao": row[7] or "",
                        "categoria": str(categoria_val)
                    })
            return produtos
        except Exception as e:
            print(f"Erro ao carregar Excel: {e}")
            return []

    def abrir_janela_edicao_planilha(self):
        produtos = self.carregar_produtos_excel()
        if not produtos:
            self.mostrar_feedback(
                "Nenhum produto cadastrado no Excel para editar!", COR_PERIGO)
            return

        janela = ctk.CTkToplevel(self)
        janela.title("Gerenciar e Editar Planilha de Produtos")
        janela.geometry("820x600")
        janela.minsize(450, 400)
        janela.configure(fg_color=COR_JANELA)
        janela.update()
        janela.grab_set()

        ctk.CTkLabel(janela, text="Produtos Salvos na Planilha",
                     font=FONTE_TITULO, text_color="white").pack(pady=15)

        frame_topo_acoes = ctk.CTkFrame(janela, fg_color="transparent")
        frame_topo_acoes.pack(fill="x", padx=20, pady=(0, 10))

        campo_pesquisa = ctk.CTkEntry(
            frame_topo_acoes, placeholder_text="Pesquisar por ID ou nome...", font=FONTE_CAMPO,
            fg_color=COR_CAMPO, border_color=COR_BORDA, width=170, height=32
        )
        campo_pesquisa.pack(side="left", padx=(0, 8))
        campo_pesquisa.delete(0, "end")

        opcoes_filtro = ["Primeiro ao último", "Último ao primeiro",
                         "Por Categoria", "Filtrar por ID (Crescente)", "Filtrar por ID (Decrescente)"]
        combo_filtro = ctk.CTkComboBox(
            frame_topo_acoes, values=opcoes_filtro, font=FONTE_LABEL, fg_color=COR_CAMPO,
            button_color=COR_BOTAO, button_hover_color=COR_BOTAO_HOVER, dropdown_fg_color=COR_CARD,
            width=160, height=32, state="readonly"
        )
        combo_filtro.set("Primeiro ao último")
        combo_filtro.pack(side="left", padx=(0, 8))

        lbl_contador = ctk.CTkLabel(
            frame_topo_acoes, text="0 selecionados", font=FONTE_LABEL, text_color=COR_AVISO)
        lbl_contador.pack(side="left", padx=5)

        checkboxes_produtos = {}

        def atualizar_contador():
            selecionados = sum(
                1 for var in checkboxes_produtos.values() if var.get() == 1)
            lbl_contador.configure(text=f"{selecionados} selecionados")

        def alternar_selecionar_todos():
            todos_marcados = all(var.get() == 1 for var in checkboxes_produtos.values(
            )) if checkboxes_produtos else False
            novo_estado = 0 if todos_marcados else 1
            for var in checkboxes_produtos.values():
                var.set(novo_estado)
            atualizar_contador()

        def excluir_selecionados():
            ids_para_excluir = [
                pid for pid, var in checkboxes_produtos.items() if var.get() == 1]
            if not ids_para_excluir:
                messagebox.showwarning(
                    "Aviso", "Nenhum produto selecionado para exclusão!",parent=janela)
                return

            if messagebox.askyesno("Confirmar Exclusão Múltipla", f"Deseja excluir os {len(ids_para_excluir)} produtos selecionados?", parent=janela):
                try:
                    wb = load_workbook(ARQUIVO_EXCEL)
                    ws = wb.active
                    if ws is not None:
                        ws = cast(Worksheet, ws)
                        linhas_para_remover = []
                        for row in range(2, ws.max_row + 1):
                            val_celula = cast(Cell, ws.cell(
                                row=row, column=1)).value
                            if val_celula is not None:
                                try:
                                    pid_val = int(float(str(val_celula)))
                                    if pid_val in ids_para_excluir:
                                        linhas_para_remover.append(row)
                                except ValueError:
                                    pass

                        for linha in sorted(linhas_para_remover, reverse=True):
                            ws.delete_rows(linha)

                        wb.save(ARQUIVO_EXCEL)
                        atualizar_lista_interface()
                        self.mostrar_feedback(
                            f"{len(linhas_para_remover)} produtos excluídos com sucesso!", COR_AVISO)
                except Exception as ex:
                    messagebox.showerror(
                        "Erro", f"Não foi possível excluir os produtos: {ex}")

        btn_excluir_selecionados = ctk.CTkButton(
            frame_topo_acoes, text="Excluir Selecionados", fg_color=COR_BOTAO, hover_color=COR_PERIGO_HOVER,
            width=130, height=32, command=excluir_selecionados
        )
        btn_excluir_selecionados.pack(side="right", padx=4)

        btn_selecionar_todos = ctk.CTkButton(
            frame_topo_acoes, text="Selecionar Todos", fg_color=COR_BOTAO, hover_color=COR_SUCESSO,
            width=110, height=32, command=alternar_selecionar_todos
        )
        btn_selecionar_todos.pack(side="right", padx=4)

        frame_lista = ctk.CTkScrollableFrame(
            janela,
            fg_color=COR_CARD,
            scrollbar_fg_color="transparent",
            scrollbar_button_color="#333333",
            scrollbar_button_hover_color="#555555"
        )
        frame_lista.pack(fill="both", expand=True, padx=20, pady=10)
        self._aplicar_efeito_autohide_scrollbar(frame_lista, COR_CARD)

        def atualizar_lista_interface(event=None):
            for widget in frame_lista.winfo_children():
                widget.destroy()

            checkboxes_produtos.clear()
            prods_atualizados = self.carregar_produtos_excel()

            if not prods_atualizados:
                janela.destroy()
                self.mostrar_feedback(
                    "Todos os produtos foram removidos da planilha.", COR_AVISO)
                return

            termo_pesquisa = (campo_pesquisa.get() or "").strip().lower()

            if termo_pesquisa:
                prods_atualizados = [
                    p for p in prods_atualizados
                    if termo_pesquisa in str(p['nome']).lower() or termo_pesquisa == str(p['id'])
                ]

            criterio_filtro = combo_filtro.get()
            if criterio_filtro == "Último ao primeiro":
                prods_atualizados = sorted(
                    prods_atualizados, key=lambda x: x['id'], reverse=True)
            elif criterio_filtro == "Por Categoria":
                prods_atualizados = sorted(prods_atualizados, key=lambda x: (
                    x['categoria'].lower() if x['categoria'] else "zzzz", x['id']))
            elif criterio_filtro == "Filtrar por ID (Crescente)":
                prods_atualizados = sorted(
                    prods_atualizados, key=lambda x: x['id'])
            elif criterio_filtro == "Filtrar por ID (Decrescente)":
                prods_atualizados = sorted(
                    prods_atualizados, key=lambda x: x['id'], reverse=True)
            else:
                prods_atualizados = sorted(
                    prods_atualizados, key=lambda x: x['id'])

            if not prods_atualizados:
                ctk.CTkLabel(
                    frame_lista, text="Nenhum produto encontrado com esse critério.", text_color="gray").pack(pady=20)
                return

            for p in prods_atualizados:
                f_row = ctk.CTkFrame(
                    frame_lista, fg_color=COR_CAMPO, corner_radius=6)
                f_row.pack(fill="x", pady=4, padx=5)

                var_check = ctk.IntVar(value=0)
                checkboxes_produtos[p['id']] = var_check

                chk = ctk.CTkCheckBox(
                    f_row, text="", variable=var_check, width=24,
                    command=atualizar_contador, fg_color=COR_PRIMARIO, hover_color=COR_PRIMARIO_HOVER
                )
                chk.pack(side="left", padx=(10, 5), pady=8)

                cat_txt = f" [{p['categoria']}]" if p['categoria'] else ""
                info_txt = f"#{p['id']} - {p['nome']} (R$ {p['preco']}){cat_txt}"
                ctk.CTkLabel(f_row, text=info_txt, font=FONTE_LABEL, text_color=COR_TEXTO, anchor="w").pack(
                    side="left", padx=5, pady=8)

                def carregar_para_edicao(prod_obj=p):
                    self.campo_nome.delete(0, "end")
                    self.campo_nome.insert(0, prod_obj["nome"])
                    self.campo_categoria.delete(0, "end")
                    self.campo_categoria.insert(0, prod_obj["categoria"])
                    self.campo_preco.delete(0, "end")
                    self.campo_preco.insert(0, prod_obj["preco"])
                    self.campo_cupom.delete(0, "end")
                    self.campo_cupom.insert(0, prod_obj["cupom"])
                    self.campo_link.delete(0, "end")
                    self.campo_link.insert(0, prod_obj["link"])

                    self.campo_link_imagem.delete(0, "end")
                    self.campo_link_imagem.insert(0, prod_obj["link_imagem"])
                    self.validar_link_imagem_tempo_real()

                    self.campo_descricao.delete("1.0", "end")
                    self.campo_descricao.insert("1.0", prod_obj["descricao"])
                    self.campo_descricao.configure(text_color=COR_TEXTO)
                    self.atualizar_descricao_preview()

                    janela.destroy()
                    self.mostrar_feedback(
                        f"Produto #{prod_obj['id']} carregado para edição!", COR_SUCESSO)

                def excluir_item(id_prod=p["id"]):
                    if messagebox.askyesno("Confirmar Exclusão", f"Deseja excluir o produto #{id_prod} da planilha?", parent=janela):
                        try:
                            wb = load_workbook(ARQUIVO_EXCEL)
                            ws = wb.active
                            if ws is not None:
                                ws = cast(Worksheet, ws)
                                linha_encontrada = None
                                for row in range(2, ws.max_row + 1):
                                    val_celula = cast(Cell, ws.cell(
                                        row=row, column=1)).value
                                    if val_celula is not None:
                                        try:
                                            if int(float(str(val_celula))) == int(id_prod):
                                                linha_encontrada = row
                                                break
                                        except ValueError:
                                            pass

                                if linha_encontrada:
                                    ws.delete_rows(linha_encontrada)
                                    wb.save(ARQUIVO_EXCEL)
                                    atualizar_lista_interface()
                                    atualizar_contador()
                                    self.mostrar_feedback(
                                        f"Produto #{id_prod} excluído com sucesso!", COR_AVISO)
                        except Exception as ex:
                            messagebox.showerror(
                                "Erro", f"Não foi possível excluir: {ex}")

                def editar_item_individual(prod_obj=p):
                    janela_editar = ctk.CTkToplevel(janela)
                    janela_editar.title(f"Editar Produto #{prod_obj['id']}")
                    janela_editar.geometry("420x520")
                    janela_editar.minsize(380, 460)
                    janela_editar.configure(fg_color=COR_JANELA)
                    janela_editar.transient(janela)
                    janela_editar.update()
                    janela_editar.grab_set()

                    ctk.CTkLabel(
                        janela_editar, text=f"Editar Produto #{prod_obj['id']}",
                        font=FONTE_TITULO, text_color="white"
                    ).pack(pady=(15, 10))

                    frame_campos_edit = ctk.CTkFrame(
                        janela_editar, fg_color="transparent")
                    frame_campos_edit.pack(fill="both", expand=True, padx=20)

                    ctk.CTkLabel(
                        frame_campos_edit, text="Nome do produto", font=FONTE_LABEL,
                        text_color=COR_TEXTO, anchor="w"
                    ).pack(fill="x", pady=(0, 2))
                    campo_nome_edit = ctk.CTkEntry(
                        frame_campos_edit, font=FONTE_CAMPO, fg_color=COR_CAMPO,
                        border_color=COR_BORDA, border_width=1, corner_radius=6,
                        text_color=COR_TEXTO, height=34
                    )
                    campo_nome_edit.pack(fill="x", pady=(0, 12))
                    campo_nome_edit.insert(0, prod_obj["nome"])

                    ctk.CTkLabel(
                        frame_campos_edit, text="Categoria", font=FONTE_LABEL,
                        text_color=COR_TEXTO, anchor="w"
                    ).pack(fill="x", pady=(0, 2))
                    campo_categoria_edit = ctk.CTkEntry(
                        frame_campos_edit, font=FONTE_CAMPO, fg_color=COR_CAMPO,
                        border_color=COR_BORDA, border_width=1, corner_radius=6,
                        text_color=COR_TEXTO, height=34
                    )
                    campo_categoria_edit.pack(fill="x", pady=(0, 12))
                    campo_categoria_edit.insert(0, prod_obj["categoria"])

                    ctk.CTkLabel(
                        frame_campos_edit, text="Preço", font=FONTE_LABEL,
                        text_color=COR_TEXTO, anchor="w"
                    ).pack(fill="x", pady=(0, 2))
                    campo_preco_edit = ctk.CTkEntry(
                        frame_campos_edit, font=FONTE_CAMPO, fg_color=COR_CAMPO,
                        border_color=COR_BORDA, border_width=1, corner_radius=6,
                        text_color=COR_TEXTO, height=34
                    )
                    campo_preco_edit.pack(fill="x", pady=(0, 12))
                    campo_preco_edit.insert(0, str(prod_obj["preco"]))

                    ctk.CTkLabel(
                        frame_campos_edit, text="Descrição", font=FONTE_LABEL,
                        text_color=COR_TEXTO, anchor="w"
                    ).pack(fill="x", pady=(0, 2))
                    campo_descricao_edit = ctk.CTkTextbox(
                        frame_campos_edit, height=120, fg_color=COR_CAMPO, border_color=COR_BORDA,
                        border_width=1, corner_radius=6, text_color=COR_TEXTO, font=FONTE_CAMPO,
                        wrap="word"
                    )
                    campo_descricao_edit.pack(
                        fill="both", expand=True, pady=(0, 10))
                    campo_descricao_edit.insert("1.0", prod_obj["descricao"])

                    def salvar_edicao_individual():
                        novo_nome = campo_nome_edit.get().strip()
                        nova_categoria = campo_categoria_edit.get().strip()
                        novo_preco = campo_preco_edit.get().strip()
                        nova_descricao = campo_descricao_edit.get(
                            "1.0", "end-1c")

                        if not novo_nome or not novo_preco:
                            messagebox.showwarning(
                                "Campos Incompletos",
                                "Nome e preço não podem ficar em branco.",
                                parent=janela_editar
                            )
                            return

                        try:
                            wb = load_workbook(ARQUIVO_EXCEL)
                            ws = wb.active
                            if ws is None:
                                raise RuntimeError("Não foi possível obter a planilha ativa.")
                            ws = cast(Worksheet, ws)
                            linha_encontrada = None
                            for row in range(2, ws.max_row + 1):
                                val_celula = cast(Cell, ws.cell(
                                    row=row, column=1)).value
                                if val_celula is not None:
                                    try:
                                        if int(float(str(val_celula))) == int(prod_obj["id"]):
                                            linha_encontrada = row
                                            break
                                    except ValueError:
                                        pass
                            if linha_encontrada:
                                cast(Cell, ws.cell(row=linha_encontrada,
                                     column=3)).value = novo_nome
                                cast(Cell, ws.cell(row=linha_encontrada,
                                     column=4)).value = novo_preco
                                cast(Cell, ws.cell(row=linha_encontrada,
                                     column=8)).value = nova_descricao
                                cast(Cell, ws.cell(row=linha_encontrada,
                                     column=9)).value = nova_categoria
                                wb.save(ARQUIVO_EXCEL)

                                janela_editar.destroy()
                                atualizar_lista_interface()
                                self.mostrar_feedback(
                                    f"Produto #{prod_obj['id']} atualizado com sucesso!", COR_SUCESSO)
                            else:
                                messagebox.showerror(
                                    "Erro", "Produto não encontrado na planilha.", parent=janela_editar)
                        except Exception as ex:
                            messagebox.showerror(
                                "Erro", f"Não foi possível salvar as alterações: {ex}", parent=janela_editar)

                    frame_botoes_edit = ctk.CTkFrame(
                        janela_editar, fg_color="transparent")
                    frame_botoes_edit.pack(fill="x", padx=20, pady=15)

                    ctk.CTkButton(
                        frame_botoes_edit, text="Cancelar", fg_color=COR_BOTAO, hover_color=COR_PERIGO_HOVER,
                        width=100, height=34, command=janela_editar.destroy
                    ).pack(side="left")

                    ctk.CTkButton(
                        frame_botoes_edit, text="Salvar Alterações", fg_color=COR_PRIMARIO, hover_color=COR_PRIMARIO_HOVER,
                        width=170, height=34, command=salvar_edicao_individual
                    ).pack(side="right")

                ctk.CTkButton(
                    f_row, text="Excluir", fg_color=COR_BOTAO, hover_color=COR_PERIGO_HOVER,
                    width=65, height=28, command=excluir_item
                ).pack(side="right", padx=8)

                ctk.CTkButton(
                    f_row, text="Carregar", fg_color=COR_BOTAO, hover_color=COR_PRIMARIO,
                    width=65, height=28, command=carregar_para_edicao
                ).pack(side="right", padx=2)

                ctk.CTkButton(
                    f_row, text="Editar", fg_color=COR_BOTAO, hover_color=COR_SUCESSO,
                    width=55, height=28, command=editar_item_individual
                ).pack(side="right", padx=2)

            atualizar_contador()

        campo_pesquisa.bind("<KeyRelease>", atualizar_lista_interface)
        combo_filtro.configure(
            command=lambda escolha: atualizar_lista_interface())

        atualizar_lista_interface()

    def enviar_para_whatsapp_grupo(self):
        janela = ctk.CTkToplevel(self)
        janela.title("Configurar Envio WhatsApp")
        janela.geometry("640x780")
        janela.minsize(450, 500)
        janela.configure(fg_color=COR_JANELA)
        janela.update()
        janela.grab_set()

        coordenadas = carregar_coordenadas()
        presets_coordenadas = carregar_presets_coordenadas()
        preset_ativo = ctk.StringVar(value="Definição 1")

        lote_selecionado = []

        ctk.CTkLabel(janela, text="Configurar Envio WhatsApp",
                     font=FONTE_TITULO, text_color="white").pack(pady=(15, 5))

        container_scroll = ctk.CTkScrollableFrame(
            janela, fg_color="transparent", scrollbar_button_color="#333333", scrollbar_button_hover_color="#555555"
        )
        container_scroll.pack(fill="both", expand=True, padx=15, pady=(0, 10))
        self._aplicar_efeito_autohide_scrollbar(container_scroll, COR_JANELA)

        frame = ctk.CTkFrame(
            container_scroll, fg_color=COR_CARD, corner_radius=10)
        frame.pack(fill="x", expand=True, padx=5, pady=5)

        cabecalho_produto = ctk.CTkFrame(frame, fg_color="transparent")
        cabecalho_produto.pack(fill="x", padx=15, pady=(15, 5))

        quadrado_cab_produto = ctk.CTkFrame(
            cabecalho_produto, width=22, height=22, fg_color=COR_QUADRADO_ICONE, corner_radius=4
        )
        quadrado_cab_produto.pack(side="left", padx=(0, 6))
        quadrado_cab_produto.pack_propagate(False)

        if self.icone_nome:
            lbl_icon_cab_produto = ctk.CTkLabel(
                quadrado_cab_produto, text="", image=self.icone_nome)
        else:
            lbl_icon_cab_produto = ctk.CTkLabel(
                quadrado_cab_produto, text="📦", font=("Segoe UI", 11))
        lbl_icon_cab_produto.place(relx=0.5, rely=0.5, anchor="center")

        ctk.CTkLabel(cabecalho_produto, text="Produto", font=(
            "Segoe UI", 11, "bold")).pack(side="left")

        f_produto_info = ctk.CTkFrame(
            frame, fg_color=COR_CAMPO, corner_radius=6)
        f_produto_info.pack(fill="x", padx=15, pady=(0, 10))

        nome_produto_atual = self.campo_nome.get().strip() or "Nenhum produto selecionado"
        lbl_info_prod = ctk.CTkLabel(
            f_produto_info, text=f"Produto: {nome_produto_atual}", font=FONTE_LABEL, text_color=COR_AVISO, anchor="w")
        lbl_info_prod.pack(side="left", padx=12, pady=10)

        def abrir_seletor_produto_cadastrado():
            produtos = self.carregar_produtos_excel()
            if not produtos:
                self.mostrar_feedback(
                    "Nenhum produto cadastrado no Excel!", COR_PERIGO)
                return

            janela_sel = ctk.CTkToplevel(janela)
            janela_sel.title("Selecionar Produto Cadastrado")
            janela_sel.geometry("480x420")
            janela_sel.minsize(360, 300)
            janela_sel.configure(fg_color=COR_JANELA)
            janela_sel.grab_set()

            ctk.CTkLabel(janela_sel, text="Escolha um Produto para Enviar",
                         font=FONTE_TITULO, text_color="white").pack(pady=15)

            f_lista = ctk.CTkScrollableFrame(
                janela_sel, fg_color=COR_CARD, scrollbar_fg_color="transparent")
            f_lista.pack(fill="both", expand=True, padx=20, pady=10)
            self._aplicar_efeito_autohide_scrollbar(f_lista, COR_CARD)

            for p in produtos:
                row_p = ctk.CTkFrame(
                    f_lista, fg_color=COR_CAMPO, corner_radius=6)
                row_p.pack(fill="x", pady=4, padx=5)

                txt_p = f"#{p['id']} - {p['nome']} (R$ {p['preco']})"
                ctk.CTkLabel(row_p, text=txt_p, font=FONTE_LABEL, text_color=COR_TEXTO, anchor="w").pack(
                    side="left", padx=10, pady=8)

                def selecionar_este(prod=p):
                    self.campo_nome.delete(0, "end")
                    self.campo_nome.insert(0, prod["nome"])
                    self.campo_categoria.delete(0, "end")
                    self.campo_categoria.insert(0, prod["categoria"])
                    self.campo_preco.delete(0, "end")
                    self.campo_preco.insert(0, prod["preco"])
                    self.campo_cupom.delete(0, "end")
                    self.campo_cupom.insert(0, prod["cupom"])
                    self.campo_link.delete(0, "end")
                    self.campo_link.insert(0, prod["link"])

                    self.campo_link_imagem.delete(0, "end")
                    self.campo_link_imagem.insert(0, prod["link_imagem"])
                    self.validar_link_imagem_tempo_real()

                    self.campo_descricao.delete("1.0", "end")
                    self.campo_descricao.insert("1.0", prod["descricao"])
                    self.campo_descricao.configure(text_color=COR_TEXTO)
                    self.atualizar_descricao_preview()

                    lote_selecionado.clear()
                    lbl_info_lote.configure(
                        text="Envio em Lote: Nenhum lote selecionado")
                    lbl_info_prod.configure(text=f"Produto: {prod['nome']}")
                    janela_sel.destroy()
                    self.mostrar_feedback(
                        f"Produto #{prod['id']} selecionado para envio!", COR_SUCESSO)

                ctk.CTkButton(
                    row_p, text="Selecionar", fg_color=COR_PRIMARIO, hover_color=COR_PRIMARIO_HOVER,
                    width=80, height=28, command=selecionar_este
                ).pack(side="right", padx=10)

        btn_usar_cadastrado = self._botao(
            f_produto_info, "Usar Cadastrado", abrir_seletor_produto_cadastrado,
            cor=COR_BOTAO, cor_hover=COR_PRIMARIO, largura=130
        )
        btn_usar_cadastrado.pack(side="right", padx=10, pady=6)

        f_lote_info = ctk.CTkFrame(frame, fg_color=COR_CAMPO, corner_radius=6)
        f_lote_info.pack(fill="x", padx=15, pady=(0, 15))

        lbl_info_lote = ctk.CTkLabel(
            f_lote_info, text="Envio em Lote: Nenhum lote selecionado", font=FONTE_LABEL, text_color=COR_AVISO, anchor="w")
        lbl_info_lote.pack(side="left", padx=12, pady=10)

        def abrir_envio_multiplos_produtos():
            produtos = self.carregar_produtos_excel()
            if not produtos:
                self.mostrar_feedback(
                    "Nenhum produto cadastrado no Excel para envio em lote!", COR_PERIGO)
                return

            janela_lote = ctk.CTkToplevel(janela)
            janela_lote.title("Selecionar Produtos para Lote")
            janela_lote.geometry("520x480")
            janela_lote.minsize(360, 300)
            janela_lote.configure(fg_color=COR_JANELA)
            janela_lote.grab_set()

            ctk.CTkLabel(janela_lote, text="Selecione os Produtos para o Lote",
                         font=FONTE_TITULO, text_color="white").pack(pady=15)

            f_lista_lote = ctk.CTkScrollableFrame(
                janela_lote, fg_color=COR_CARD, scrollbar_fg_color="transparent")
            f_lista_lote.pack(fill="both", expand=True, padx=20, pady=10)
            self._aplicar_efeito_autohide_scrollbar(f_lista_lote, COR_CARD)

            checks_lote = {}
            for p in produtos:
                row_l = ctk.CTkFrame(
                    f_lista_lote, fg_color=COR_CAMPO, corner_radius=6)
                row_l.pack(fill="x", pady=4, padx=5)

                var_l = ctk.IntVar(value=0)
                checks_lote[p['id']] = (var_l, p)

                chk_l = ctk.CTkCheckBox(
                    row_l, text="", variable=var_l, width=24, fg_color=COR_PRIMARIO)
                chk_l.pack(side="left", padx=(10, 5), pady=8)

                txt_l = f"#{p['id']} - {p['nome']} (R$ {p['preco']})"
                ctk.CTkLabel(row_l, text=txt_l, font=FONTE_LABEL, text_color=COR_TEXTO, anchor="w").pack(
                    side="left", padx=5, pady=8)

            def confirmar_selecao_lote():
                selecionados = [
                    p_obj for pid, (v, p_obj) in checks_lote.items() if v.get() == 1]
                if not selecionados:
                    messagebox.showwarning(
                        "Aviso", "Selecione ao menos um produto para o lote!", parent=janela_lote)
                    return

                lote_selecionado.clear()
                lote_selecionado.extend(selecionados)

                lbl_info_lote.configure(
                    text=f"Envio em Lote: {len(lote_selecionado)} produtos selecionados")
                janela_lote.destroy()
                self.mostrar_feedback(
                    f"Lote de {len(lote_selecionado)} produtos selecionado com sucesso!", COR_SUCESSO)

            self._botao(janela_lote, "Confirmar Seleção de Lote", confirmar_selecao_lote,
                        cor=COR_SUCESSO, cor_hover="#1ebe5d", largura=200).pack(pady=15)

        self._botao(f_lote_info, "Selecionar Lote", abrir_envio_multiplos_produtos,
                    cor=COR_BOTAO, cor_hover=COR_WHATSAPP, largura=130).pack(side="right", padx=10, pady=6)

        cabecalho_agendar = ctk.CTkFrame(frame, fg_color="transparent")
        cabecalho_agendar.pack(fill="x", padx=15, pady=(15, 5))

        quadrado_agendar = ctk.CTkFrame(
            cabecalho_agendar, width=22, height=22, fg_color=COR_QUADRADO_ICONE, corner_radius=4
        )
        quadrado_agendar.pack(side="left", padx=(0, 6))
        quadrado_agendar.pack_propagate(False)

        if self.icone_agendar:
            lbl_icon_agendar = ctk.CTkLabel(
                quadrado_agendar, text="", image=self.icone_agendar)
        else:
            lbl_icon_agendar = ctk.CTkLabel(
                quadrado_agendar, text="⏰", font=("Segoe UI", 11))
        lbl_icon_agendar.place(relx=0.5, rely=0.5, anchor="center")

        ctk.CTkLabel(cabecalho_agendar, text="Agendar Horário de Envio (Opcional)", font=(
            "Segoe UI", 11, "bold")).pack(side="left")

        f_agendar = ctk.CTkFrame(frame, fg_color=COR_CAMPO, corner_radius=6)
        f_agendar.pack(fill="x", padx=15, pady=(0, 15))

        var_agendar = ctk.IntVar(value=0)
        chk_agendar = ctk.CTkCheckBox(f_agendar, text="Ativar Agendamento",
                                      variable=var_agendar, font=FONTE_LABEL, fg_color=COR_PRIMARIO)
        chk_agendar.pack(side="left", padx=12, pady=10)

        campo_horario = ctk.CTkEntry(f_agendar, placeholder_text="Ex: 14:30 ou 14:30:00",
                                     width=150, height=32, fg_color=COR_CARD, border_color=COR_BORDA)
        campo_horario.pack(side="right", padx=12, pady=10)

        def evento_horario(e):
            self.validar_horario(e)
        campo_horario.bind("<KeyRelease>", evento_horario)
        campo_horario.bind("<FocusOut>", self.atualizar_horario_campo)

        ctk.CTkLabel(f_agendar, text="Horário (HH:MM):", font=FONTE_LABEL,
                     text_color="gray").pack(side="right", padx=(0, 5))

        cabecalho_destino = ctk.CTkFrame(frame, fg_color="transparent")
        cabecalho_destino.pack(fill="x", padx=15, pady=(15, 5))

        quadrado_destino = ctk.CTkFrame(
            cabecalho_destino, width=22, height=22, fg_color=COR_QUADRADO_ICONE, corner_radius=4
        )
        quadrado_destino.pack(side="left", padx=(0, 6))
        quadrado_destino.pack_propagate(False)

        if self.icone_destino:
            lbl_icon_destino = ctk.CTkLabel(
                quadrado_destino, text="", image=self.icone_destino)
        else:
            lbl_icon_destino = ctk.CTkLabel(
                quadrado_destino, text="🎯", font=("Segoe UI", 11))
        lbl_icon_destino.place(relx=0.5, rely=0.5, anchor="center")

        ctk.CTkLabel(cabecalho_destino, text="Destino", font=(
            "Segoe UI", 11, "bold")).pack(side="left")

        f_grupo_wrapper = ctk.CTkFrame(frame, fg_color="transparent")
        f_grupo_wrapper.pack(fill="x", padx=15, pady=(0, 15))

        ctk.CTkLabel(f_grupo_wrapper, text="Nome do Contato/Grupo", font=FONTE_LABEL,
                     text_color="gray").pack(anchor="w", pady=(0, 2))
        campo_grupo = ctk.CTkEntry(f_grupo_wrapper, placeholder_text="Ex: Ofertas do Dia",
                                   height=35, fg_color=COR_CAMPO, border_color=COR_BORDA)
        campo_grupo.pack(fill="x")

        f_instrucoes = ctk.CTkFrame(
            frame, fg_color="#1d222a", border_color=COR_PRIMARIO, border_width=1, corner_radius=6)
        f_instrucoes.pack(fill="x", padx=15, pady=(5, 10))

        lbl_instrucoes_titulo = ctk.CTkLabel(f_instrucoes, text="Como configurar as coordenadas para funcionar:", font=(
            "Segoe UI", 11, "bold"), text_color=COR_TEXTO)
        lbl_instrucoes_titulo.pack(anchor="w", padx=10, pady=(6, 2))

        texto_instrucoes = (
            "1. Deixe a janela do WhatsApp Web visível na tela.\n"
            "2. Clique em 'Capturar' e, em 3 segundos, posicione o mouse no local correspondente no WhatsApp.\n"
            "3. Use os botões 'Definição 1, 2 ou 3' para carregar posições salvas e clique em 'Salvar na Definição'."
        )
        lbl_instrucoes_corpo = ctk.CTkLabel(f_instrucoes, text=texto_instrucoes, font=(
            "Segoe UI", 10), text_color="gray", justify="left")
        lbl_instrucoes_corpo.pack(anchor="w", padx=10, pady=(0, 6))

        f_presets = ctk.CTkFrame(frame, fg_color="transparent")
        f_presets.pack(fill="x", padx=15, pady=(5, 5))

        ctk.CTkLabel(f_presets, text="Definições Salvas:",
                     font=("Segoe UI", 11, "bold")).pack(side="left")

        campos_coord = {}
        botoes_def_dict = {}

        def atualizar_labels_coords():
            for nome_key, lbl_widget in campos_coord.items():
                c = coordenadas.get(nome_key, {"x": 0, "y": 0})
                lbl_widget.configure(text=f"{c['x']}, {c['y']}")

        def carregar_preset_selecionado(nome_def):
            preset_ativo.set(nome_def)
            if nome_def in presets_coordenadas:
                coordenadas.update(presets_coordenadas[nome_def])
                salvar_coordenadas(coordenadas)
                atualizar_labels_coords()
                lbl_status_preset.configure(
                    text=f"Carregado: {nome_def}", text_color=COR_SUCESSO)

            for b_nome, btn_obj in botoes_def_dict.items():
                if b_nome == nome_def:
                    btn_obj.configure(border_width=1, border_color=COR_SUCESSO)
                else:
                    btn_obj.configure(border_width=0)

        def salvar_preset_atual():
            nome_def = preset_ativo.get()
            presets_coordenadas[nome_def] = {k: v.copy() if isinstance(
                v, dict) else v for k, v in coordenadas.items()}
            salvar_presets_coordenadas(presets_coordenadas)
            salvar_coordenadas(coordenadas)
            lbl_status_preset.configure(
                text=f"Salvo em '{nome_def}' com sucesso!", text_color=COR_SUCESSO)

        for nome_def in ["Definição 1", "Definição 2", "Definição 3"]:
            btn_def = ctk.CTkButton(
                f_presets, text=nome_def, width=75, height=26, font=("Segoe UI", 10),
                fg_color=COR_BOTAO, hover_color=COR_BOTAO_HOVER,
                command=lambda d=nome_def: carregar_preset_selecionado(d)
            )
            btn_def.pack(side="left", padx=3)
            botoes_def_dict[nome_def] = btn_def

        if "Definição 1" in botoes_def_dict:
            botoes_def_dict["Definição 1"].configure(
                border_width=1, border_color=COR_SUCESSO)

        btn_salvar_def = ctk.CTkButton(
            f_presets, text="Salvar na Definição", width=110, height=26, font=("Segoe UI", 10),
            fg_color=COR_BOTAO, hover_color=COR_PRIMARIO_HOVER, command=salvar_preset_atual
        )
        btn_salvar_def.pack(side="right")

        lbl_status_preset = ctk.CTkLabel(
            frame, text="Definição 1 Selecionada", font=("Segoe UI", 10), text_color="gray")
        lbl_status_preset.pack(anchor="w", padx=15, pady=(0, 5))

        itens = [("grupo", "Grupo/Contato"), ("mensagem", "Campo Mensagem"),
                 ("anexo", "Botão Anexo"), ("fotos", "Fotos e Vídeos")]

        for nome, titulo in itens:
            f_item = ctk.CTkFrame(frame, fg_color="transparent")
            f_item.pack(fill="x", padx=15, pady=3)
            ctk.CTkLabel(f_item, text=titulo, width=120,
                         anchor="w").pack(side="left")

            coord = coordenadas.get(nome, {"x": 0, "y": 0})
            lbl = ctk.CTkLabel(
                f_item, text=f"{coord['x']}, {coord['y']}", text_color=COR_SUCESSO, width=80)
            lbl.pack(side="left")
            campos_coord[nome] = lbl

            def acao_captura(l=lbl, n=nome):
                def cap():
                    l.configure(text="Aguarde 3s...", text_color=COR_AVISO)
                    janela.after(
                        3000, lambda: self._finalizar_captura_coord(l, n, coordenadas))
                return cap

            self._botao(f_item, "Capturar", acao_captura(),
                        largura=75).pack(side="right")

            canvas_linha = ctk.CTkCanvas(
                f_item, height=10, bg=COR_CARD, highlightthickness=0, bd=0)
            canvas_linha.pack(side="left", fill="x", expand=True, padx=8)

            def desenhar_linha_tracejada(event, c=canvas_linha):
                c.delete("all")
                c.create_line(0, event.height // 2, event.width,
                              event.height // 2, dash=(4, 4), fill="#555555")

            canvas_linha.bind("<Configure>", desenhar_linha_tracejada)

        def rotina_disparo(produtos_envio, nome_g, agendamento_obj=None):
            self.cancel_event.clear()
            if agendamento_obj:
                agendamento_obj["status"] = "Em andamento"

            sucesso_geral = True
            for idx, prod_item in enumerate(produtos_envio, 1):
                if self.cancel_event.is_set():
                    self.mostrar_feedback(
                        "Envio cancelado pelo usuário!", COR_PERIGO)
                    if agendamento_obj:
                        agendamento_obj["status"] = "Cancelado"
                    return

                msg = prod_item.get("descricao", "")
                img = prod_item.get("link_imagem", "").strip() or None

                self.mostrar_feedback(
                    f"Enviando item {idx} de {len(produtos_envio)}...", COR_AVISO)
                sucesso, res_msg = enviar_para_grupo(
                    nome_g, msg, img, cancel_event=self.cancel_event)
                if not sucesso:
                    sucesso_geral = False
                    if res_msg == 'Envio cancelado pelo usuário.':
                        if agendamento_obj:
                            agendamento_obj["status"] = "Cancelado"
                        return

            if sucesso_geral:
                self.mostrar_feedback(
                    "Envio(s) concluído(s) com sucesso!", COR_SUCESSO)
                if agendamento_obj:
                    agendamento_obj["status"] = "Concluído"
            else:
                self.mostrar_feedback(
                    "Envio(s) finalizado(s) com avisos.", COR_AVISO)
                if agendamento_obj:
                    agendamento_obj["status"] = "Concluído c/ erros"

        def processar_envio():
            nome_prod = self.campo_nome.get().strip()
            if not lote_selecionado and not nome_prod:
                messagebox.showwarning(
                    "Aviso",
                    "Selecione um produto cadastrado ou um lote de produtos antes de prosseguir com o envio.",
                    parent=janela
                )
                return

            nome_g = campo_grupo.get().strip()
            if not nome_g:
                messagebox.showwarning(
                    "Aviso", "Informe o Nome do Contato/Grupo!", parent=janela)
                return

            salvar_coordenadas(coordenadas)

            if lote_selecionado:
                lista_disparo = lote_selecionado
            else:
                desc_atual = self.campo_descricao.get("1.0", "end-1c")
                if desc_atual == "Escreva ou aplique um template para gerar a descrição...":
                    desc_atual = ""
                img_atual = self.campo_link_imagem.get().strip() or None
                lista_disparo = [{
                    "nome": self.campo_nome.get().strip() or "Produto",
                    "descricao": desc_atual,
                    "link_imagem": img_atual or ""
                }]

            if var_agendar.get() == 1:
                horario_txt = campo_horario.get().strip()
                if not horario_txt:
                    messagebox.showwarning(
                        "Aviso", "Digite o horário do envio agendado (ex: 14:30)!", parent=janela)
                    return
                try:
                    partes = [int(p) for p in horario_txt.split(":")]
                    agora = datetime.now()
                    if len(partes) == 2:
                        alvo = agora.replace(
                            hour=partes[0], minute=partes[1], second=0, microsecond=0)
                    elif len(partes) == 3:
                        alvo = agora.replace(
                            hour=partes[0], minute=partes[1], second=partes[2], microsecond=0)
                    else:
                        raise ValueError

                    if alvo <= agora:
                        alvo += timedelta(days=1)

                    atraso_segundos = (alvo - agora).total_seconds()

                    agendamento_id = len(self.agendamentos) + 1
                    novo_agendamento = {
                        "id": agendamento_id,
                        "horario_str": alvo.strftime("%d/%m/%Y às %H:%M:%S"),
                        "horario_dt": alvo,
                        "grupo": nome_g,
                        "produtos": lista_disparo,
                        "status": "Agendado",
                        "timer": None
                    }

                    def tarefa_agendada():
                        rotina_disparo(lista_disparo, nome_g, novo_agendamento)

                    timer = threading.Timer(atraso_segundos, tarefa_agendada)
                    novo_agendamento["timer"] = timer
                    self.agendamentos.append(novo_agendamento)
                    timer.daemon = True
                    timer.start()

                    janela.destroy()
                    self.mostrar_feedback(
                        f"Agendamento #{agendamento_id} programado para {alvo.strftime('%d/%m/%Y às %H:%M:%S')}!", COR_SUCESSO)

                except ValueError:
                    messagebox.showerror(
                        "Formato Inválido", "Informe um horário válido em HH:MM ou HH:MM:SS.", parent=janela)
            else:
                janela.destroy()
                self.mostrar_feedback(
                    "Iniciando processo de envio...", COR_AVISO)
                t = threading.Thread(target=rotina_disparo,
                                     args=(lista_disparo, nome_g))
                t.daemon = True
                t.start()

        self._botao(janela, "Confirmar e Executar Envio", processar_envio,
                    cor=COR_BOTAO, cor_hover="#1ebe5d", largura=220).pack(pady=12)

    def _finalizar_captura_coord(self, label, nome, coordenadas):
        pos = pyautogui.position()
        coordenadas[nome] = {"x": pos[0], "y": pos[1]}
        label.configure(text=f"{pos[0]}, {pos[1]}", text_color=COR_SUCESSO)


if __name__ == "__main__":
    app = AplicacaoAfiliado()
    app.mainloop()
